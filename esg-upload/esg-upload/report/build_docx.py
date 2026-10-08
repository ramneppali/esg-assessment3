"""Convert Assessment3_Report.md into a Word document (.docx) using python-docx.

Usage: python3 build_docx.py            (full report)
       python3 build_docx.py --short    (condensed copy, full report left as is)
(needs: pip install python-docx)
Handles headings, paragraphs, bullets, numbered lists, blockquotes, tables,
fenced code blocks and **bold** / *italic* / `code` inline text.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Inches, Pt
except ImportError:
    sys.exit("Install python-docx first: pip install python-docx")

HERE = Path(__file__).resolve().parent
SRC = HERE / "Assessment3_Report.md"
OUT = HERE / "Assessment3_Report.docx"
OUT_SHORT = HERE / "Assessment3_Report_Short.docx"

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")
IMAGE = re.compile(r"!\[(.*?)\]\((.+?)\)")


def add_image(doc, caption: str, path: Path, max_w: float = 6.0, max_h: float = 8.0) -> None:
    if not path.is_file():
        doc.add_paragraph(f"[Missing image: {path.name}]")
        return
    width = Inches(max_w)
    try:
        from PIL import Image
        with Image.open(path) as img:
            w_px, h_px = img.size
        # Keep tall images within the height limit.
        width = Inches(min(max_w, max_h * w_px / h_px))
    except (ImportError, OSError, ZeroDivisionError):
        pass
    doc.add_picture(str(path), width=width)
    par = doc.add_paragraph()
    par.add_run(caption).italic = True


def add_inline(par, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Courier New"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows: list[list[str]], font_pt: float | None = None) -> None:
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            add_inline(cell.paragraphs[0], row[j] if j < len(row) else "")
            for run in cell.paragraphs[0].runs:
                if i == 0:
                    run.bold = True
                if font_pt:
                    run.font.size = Pt(font_pt)
    doc.add_paragraph()


def build(short: bool = False) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10 if short else 11)
    img_w, img_h, table_pt = (4.8, 4.6, 8.5) if short else (6.0, 8.0, None)
    if short:
        for section in doc.sections:
            section.left_margin = section.right_margin = Inches(0.8)
            section.top_margin = section.bottom_margin = Inches(0.8)

    text = SRC.read_text(encoding="utf-8")
    if short:
        from shorten import shorten
        text = shorten(text)
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            par = doc.add_paragraph()
            run = par.add_run("\n".join(block))
            run.font.name = "Courier New"
            run.font.size = Pt(7.5)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.fullmatch(r"\|[\s:|-]+\|?", lines[i]):
                    rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows, table_pt)
            continue
        elif re.match(r"#{1,4} ", line):
            level = len(line.split(" ", 1)[0])
            doc.add_heading(line[level + 1:].strip(), level=min(level, 4))
        elif IMAGE.fullmatch(line.strip()):
            caption, rel = IMAGE.fullmatch(line.strip()).groups()
            add_image(doc, caption, (HERE / rel).resolve(), img_w, img_h)
        elif line.strip() == "---":
            pass
        elif line.startswith(">"):
            par = doc.add_paragraph(style="Intense Quote")
            add_inline(par, line.lstrip("> ").strip())
        elif re.match(r"\s*- ", line):
            add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*- ", "", line))
        elif re.match(r"\d+\. ", line):
            add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        elif line.strip():
            add_inline(doc.add_paragraph(), line.strip())
        i += 1

    out = OUT_SHORT if short else OUT
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    build(short="--short" in sys.argv)
