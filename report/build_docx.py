"""Convert Assessment3_Report.md into a Word document (.docx) using python-docx.

Usage: python3 build_docx.py   (needs: pip install python-docx)
Handles headings, paragraphs, bullets, numbered lists, blockquotes, tables,
fenced code blocks and **bold** / *italic* / `code` inline text.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    sys.exit("Install python-docx first: pip install python-docx")

HERE = Path(__file__).resolve().parent
SRC = HERE / "Assessment3_Report.md"
OUT = HERE / "Assessment3_Report.docx"

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


def add_inline(par, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Courier New"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows: list[list[str]]) -> None:
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            add_inline(cell.paragraphs[0], row[j] if j < len(row) else "")
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            par = doc.add_paragraph()
            run = par.add_run("\n".join(block))
            run.font.name = "Courier New"
            run.font.size = Pt(7.5)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.fullmatch(r"\|[\s:|-]+\|?", lines[i]):
                    rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        elif re.match(r"#{1,4} ", line):
            level = len(line.split(" ", 1)[0])
            doc.add_heading(line[level + 1:].strip(), level=min(level, 4))
        elif line.strip() == "---":
            pass
        elif line.startswith(">"):
            par = doc.add_paragraph(style="Intense Quote")
            add_inline(par, line.lstrip("> ").strip())
        elif re.match(r"\s*- ", line):
            add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*- ", "", line))
        elif re.match(r"\d+\. ", line):
            add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        elif line.strip():
            add_inline(doc.add_paragraph(), line.strip())
        i += 1

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
